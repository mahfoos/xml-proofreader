"""Tests for reading the style guide .docx into text."""

import zipfile
from pathlib import Path

import docx
import pytest
from docx.enum.style import WD_STYLE_TYPE

from xml_proofreader.style_guide import StyleGuideError, load_style_guide

FIXTURES = Path(__file__).parent / "fixtures"
GUIDE = load_style_guide(FIXTURES / "style_guide.docx")
LINES = GUIDE.split("\n")


# --- The assignment's style guide


def test_custom_heading_styles_become_markdown_headings() -> None:
    assert "## Dates (Months, Days, Years)" in LINES
    assert "### $ and other currency symbols/designations" in LINES


def test_fully_bold_paragraphs_become_bold_labels() -> None:
    assert "**List of Preferred Word Usage**" in LINES


def test_list_paragraphs_become_bullets() -> None:
    assert any(line.startswith("- “To Be” standing alone is a verb") for line in LINES)


def test_rule_lines_survive_word_for_word() -> None:
    for rule in [
        "January 2022 not January, 2022",
        "wholly owned subsidiary not wholly-owned subsidiary",
        "HIPAA (not HIPPA)",
        "versus (lowercase in text; capitalized in titles and headings; never “vs.”)",
    ]:
        assert rule in LINES


def test_whitespace_is_trimmed_and_blank_paragraphs_are_dropped() -> None:
    assert "$3,000 or US$3,000 or USD3,000." in LINES
    assert all(line == line.strip() for line in LINES)
    assert "\n\n\n" not in GUIDE


# --- Any other style guide


def test_built_in_styles_and_tables_are_rendered(tmp_path: Path) -> None:
    document = docx.Document()
    document.add_heading("House Style", level=0)
    document.add_heading("Dates", level=1)
    document.add_paragraph("No comma between month and year.", style="List Bullet")
    document.add_paragraph().add_run("Preferred words").bold = True
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "email"
    table.cell(0, 1).text = "lowercase, no hyphen"
    table.cell(1, 0).text = "website"
    table.cell(1, 1).text = "one word"
    path = tmp_path / "guide.docx"
    document.save(str(path))

    assert load_style_guide(path) == (
        "# House Style\n"
        "\n"
        "## Dates\n"
        "- No comma between month and year.\n"
        "**Preferred words**\n"
        "email | lowercase, no hyphen\n"
        "website | one word"
    )


def test_a_missing_file_raises_a_clear_error(tmp_path: Path) -> None:
    with pytest.raises(StyleGuideError, match="Cannot read"):
        load_style_guide(tmp_path / "missing.docx")


def test_a_file_that_is_not_a_word_document_raises_a_clear_error(tmp_path: Path) -> None:
    path = tmp_path / "notes.docx"
    path.write_text("just some text")
    with pytest.raises(StyleGuideError, match="Cannot read"):
        load_style_guide(path)


def test_a_word_template_raises_a_clear_error(tmp_path: Path) -> None:
    document_path = tmp_path / "guide.docx"
    docx.Document().save(str(document_path))
    # A .dotx template is a .docx package whose main part has a different content type.
    template_path = tmp_path / "guide.dotx"
    with zipfile.ZipFile(document_path) as source, zipfile.ZipFile(template_path, "w") as target:
        for item in source.infolist():
            data = source.read(item)
            if item.filename == "[Content_Types].xml":
                data = data.replace(b"document.main+xml", b"template.main+xml")
            target.writestr(item, data)
    with pytest.raises(StyleGuideError, match="Cannot read"):
        load_style_guide(template_path)


def test_an_empty_document_raises_a_clear_error(tmp_path: Path) -> None:
    path = tmp_path / "empty.docx"
    docx.Document().save(str(path))
    with pytest.raises(StyleGuideError, match="no text"):
        load_style_guide(path)


def test_a_style_without_a_name_is_treated_as_plain_text(tmp_path: Path) -> None:
    document = docx.Document()
    style = document.styles.add_style("Temporary", WD_STYLE_TYPE.PARAGRAPH)
    document.add_paragraph("Rule text", style="Temporary")
    style.name = None
    path = tmp_path / "guide.docx"
    document.save(str(path))
    assert load_style_guide(path) == "Rule text"
